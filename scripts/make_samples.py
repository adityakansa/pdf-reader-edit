#!/usr/bin/env python3
"""FR-020: builds the four bundled "Sample File" documents in app/src/main/assets/samples/.

Each one exercises what its viewer claims to support, so a fresh install shows the app working:
the PDF has a real text layer (search, highlight, AI), the DOCX has headings, runs, bullets, a table and an
image, the XLSX has two sheets with shared strings and numbers, and the PPTX has text and picture slides.

Requires: pip install reportlab python-docx openpyxl python-pptx pillow
"""
import io
import os

from PIL import Image, ImageDraw

OUT = os.path.join(os.path.dirname(__file__), "..", "app", "src", "main", "assets", "samples")

BRAND_RED = (211, 47, 47)
TITLE = "Welcome to PDF Reader"
INTRO = (
    "This sample file ships with the app so every list has something to open. It disappears as soon "
    "as a real document of the same type is on your phone."
)
FEATURES = [
    ("Read", "Open PDF, Word, Excel and PowerPoint files from anywhere on your phone."),
    ("Search", "Find any word inside a PDF and jump between the matches."),
    ("Highlight", "Mark important sentences; the highlights open in any other PDF app."),
    ("Create", "Turn photos into a PDF or scan paper pages with automatic edge detection."),
    ("Sign", "Draw your signature once and place it on any PDF page."),
    ("AI Translate", "Translate a page into Italian, Polish, Portuguese, Dutch, French, German or Spanish."),
    ("AI Summary", "Get a short summary of a page, written on your phone without an internet connection."),
]
BODY = (
    "Everything happens on your device. Documents are never uploaded: translation and summaries run "
    "on the phone itself, and the files you create are saved to Documents/PDF Reader so any file app "
    "can see them.\n\n"
    "Scanning works best on a flat, well-lit surface with some contrast between the page and the "
    "table. Hold the phone still and the page is captured automatically once its edges are steady. "
    "You can then adjust the corners, rotate the page and choose a filter before saving.\n\n"
    "Signatures you draw are stored only on this phone, as transparent images, and the original PDF "
    "is never changed: a signed copy is saved beside your other outputs."
)


def banner_png(width=800, height=260):
    image = Image.new("RGB", (width, height), (246, 246, 246))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((40, 40, 220, 220), radius=28, fill=BRAND_RED)
    draw.polygon([(170, 40), (220, 90), (170, 90)], fill=(255, 205, 210))
    draw.rectangle((75, 120, 185, 132), fill="white")
    draw.rectangle((75, 150, 185, 162), fill="white")
    draw.rectangle((75, 180, 150, 192), fill="white")
    for i, colour in enumerate([(229, 57, 53), (30, 111, 217), (30, 158, 90), (244, 115, 31)]):
        x = 280 + i * 125
        draw.rounded_rectangle((x, 90, x + 100, 190), radius=18, fill=colour)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue()


def make_pdf(path):
    from reportlab.lib.colors import Color
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image as RLImage, PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    red = Color(*[c / 255 for c in BRAND_RED])
    title = ParagraphStyle("t", parent=styles["Title"], textColor=red)
    body = ParagraphStyle("b", parent=styles["BodyText"], fontSize=11, leading=16)
    story = [
        Paragraph(TITLE, title),
        RLImage(io.BytesIO(banner_png()), width=160 * mm, height=52 * mm),
        Spacer(1, 6 * mm),
        Paragraph(INTRO, body),
        Spacer(1, 4 * mm),
    ]
    for name, text in FEATURES:
        story.append(Paragraph(f"<b>{name}.</b> {text}", body))
    story.append(PageBreak())
    story.append(Paragraph("How it works", styles["Heading1"]))
    for paragraph in BODY.split("\n\n"):
        story.append(Paragraph(paragraph, body))
        story.append(Spacer(1, 3 * mm))
    SimpleDocTemplate(path, pagesize=A4, title="Sample File", author="PDF Reader").build(story)


def make_docx(path):
    from docx import Document
    from docx.shared import Mm, RGBColor

    document = Document()
    document.core_properties.title = "Sample File"
    heading = document.add_heading(TITLE, level=1)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(*BRAND_RED)
    document.add_picture(io.BytesIO(banner_png()), width=Mm(150))
    intro = document.add_paragraph()
    intro.add_run("This is a ").italic = False
    intro.add_run("sample Word document").bold = True
    intro.add_run(". " + INTRO.split(". ", 1)[1])
    document.add_heading("What you can do", level=2)
    for name, text in FEATURES:
        item = document.add_paragraph(style="List Bullet")
        item.add_run(name + ": ").bold = True
        item.add_run(text)
    document.add_heading("File types", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["Type", "Extensions", "Badge"]):
        cell.text = label
    for row in [("PDF", "pdf", "red"), ("Word", "doc, docx", "blue"),
                ("Excel", "xls, xlsx, csv", "green"), ("PowerPoint", "ppt, pptx", "orange")]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    document.add_heading("How it works", level=2)
    for paragraph in BODY.split("\n\n"):
        p = document.add_paragraph(paragraph)
    p.add_run(" Underlined text shows up too.").underline = True
    document.save(path)


def make_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monthly budget"
    sheet.append(["Item", "Category", "January", "February", "March"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    rows = [
        ("Rent", "Home", 950, 950, 950),
        ("Groceries", "Food", 310.5, 288.2, 325.75),
        ("Electricity", "Home", 64.3, 58.9, 51.2),
        ("Internet", "Home", 29.99, 29.99, 29.99),
        ("Transport", "Travel", 85, 92.4, 78),
        ("Eating out", "Food", 120, 95.5, 140.25),
        ("Books", "Leisure", 24.99, 0, 18.5),
    ]
    for row in rows:
        sheet.append(row)
    # Numbers "as stored": the viewer does not evaluate formulas, so totals are written as values.
    sheet.append(["Total", ""] + [round(sum(r[i] for r in rows), 2) for i in range(2, 5)])
    sheet.column_dimensions["A"].width = 16

    tasks = workbook.create_sheet("Tasks")
    tasks.append(["Task", "Owner", "Status", "Due"])
    for cell in tasks[1]:
        cell.font = Font(bold=True)
    for row in [
        ("Scan receipts", "Alex", "Done", "2023-01-05"),
        ("Sign lease", "Sam", "In progress", "2023-01-12"),
        ("Translate letter", "Alex", "To do", "2023-01-20"),
        ("Summarise report", "Kim", "To do", "2023-01-31"),
    ]:
        tasks.append(row)
    workbook.properties.title = "Sample File"
    workbook.save(path)


def make_pptx(path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.core_properties.title = "Sample File"

    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = TITLE
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(*BRAND_RED)
    slide.placeholders[1].text = "A sample presentation"

    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "What you can do"
    frame = slide.placeholders[1].text_frame
    frame.text = FEATURES[0][0] + " — " + FEATURES[0][1]
    for name, text in FEATURES[1:5]:
        paragraph = frame.add_paragraph()
        paragraph.text = f"{name} — {text}"
        paragraph.font.size = Pt(18)

    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Every file type in one place"
    slide.shapes.add_picture(io.BytesIO(banner_png()), Inches(0.75), Inches(2.2), width=Inches(8.5))

    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "On-device AI"
    frame = slide.placeholders[1].text_frame
    frame.text = FEATURES[5][1]
    frame.add_paragraph().text = FEATURES[6][1]
    frame.add_paragraph().text = "No page text ever leaves your phone."
    deck.save(path)


def main():
    os.makedirs(OUT, exist_ok=True)
    for name, maker in [("sample.pdf", make_pdf), ("sample.docx", make_docx),
                        ("sample.xlsx", make_xlsx), ("sample.pptx", make_pptx)]:
        target = os.path.join(OUT, name)
        maker(target)
        print(f"{name}: {os.path.getsize(target):,} bytes")


if __name__ == "__main__":
    main()
